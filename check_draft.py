import docx

doc = docx.Document("artikel/Draft Proposal TA - Copy (1).docx")

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print("="*60)
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if text:
        print(f"P[{i:02d}] Style: {p.style.name:12s} | Text: {text}")
print("="*60)
